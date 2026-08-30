from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from goblin_spatial.reporting import (
    add_sc3_reporting_aliases,
    build_sc3_land_accounting_summary,
    build_sc3_validation_table,
    export_scientific_results,
)


def _sc3_frame() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "CSOED": ["1", "2"],
            "SC2_POTENTIAL_RELEASE_HA": [100.0, 100.0],
            "SC3_STAGE_A_REALIZED_HA": [60.0, 50.0],
            "SC3_STAGE_A_AVAILABLE_HA": [40.0, 50.0],
            "SC3_REALIZED_REWETTING_HA": [10.0, 15.0],
            "SC3_RESIDUAL_AVAILABLE_LAND_HA": [30.0, 35.0],
        }
    )


def test_sc3_reporting_distinguishes_parent_post_stage_and_final_residual() -> None:
    frame = add_sc3_reporting_aliases(_sc3_frame())
    summary = build_sc3_land_accounting_summary(
        frame,
        gross_release_ha=200.0,
        stage_a_target_ha=120.0,
        parent_available_target_ha=80.0,
        rewetting_target_ha=30.0,
    ).iloc[0]

    assert np.isclose(summary["GOBLIN_PARENT_AVAILABLE_TARGET_HA"], 80.0)
    assert np.isclose(summary["STAGE_A_REALIZED_HA"], 110.0)
    assert np.isclose(summary["STAGE_A_UNMET_HA"], 10.0)
    assert np.isclose(
        summary["SC3_POST_STAGE_A_UNALLOCATED_RELEASE_HA"],
        90.0,
    )
    assert np.isclose(summary["POST_STAGE_A_MINUS_PARENT_TARGET_HA"], 10.0)
    assert np.isclose(summary["REWETTING_REALIZED_HA"], 25.0)
    assert np.isclose(summary["REWETTING_UNMET_HA"], 5.0)
    assert np.isclose(
        summary["SC3_FINAL_UNALLOCATED_RELEASED_LAND_HA"],
        65.0,
    )
    assert summary["PARENT_TARGET_ACCOUNTING_STATUS"] == "PASS"
    assert summary["SPATIAL_STAGE_A_ACCOUNTING_STATUS"] == "PASS"
    assert summary["STRICT_SPATIAL_ACCOUNTING_STATUS"] == "PASS"

    validation = build_sc3_validation_table(frame, pd.DataFrame([summary]))
    assert set(validation["STATUS"]) == {"PASS"}


def test_scientific_workbook_export_preserves_canonical_csvs(tmp_path: Path) -> None:
    run_dir = tmp_path / "SI_SG_2020_PRORATA"
    run_dir.mkdir()

    pd.DataFrame(
        [
            {
                "SCENARIO_NO": 1,
                "SCENARIO_ID": "SI_SG",
                "SCENARIO_NAME": "SI split gas",
                "RUN_START_YEAR": 2020,
                "TARGET_YEAR": 2050,
                "ALLOCATION_POLICY": "PRORATA",
                "PROTECTION_STRENGTH_LAMBDA": 0.5,
                "BASELINE_GRASSLAND_HA": 300.0,
                "TARGET_LIVESTOCK_LAND_HA": 100.0,
                "RUN_GROSS_RELEASE_HA": 200.0,
                "STAGE_A_TARGET_HA": 120.0,
                "GOBLIN_PARENT_AVAILABLE_BEFORE_REWETTING_HA": 80.0,
                "REWETTING_TARGET_HA": 30.0,
            }
        ]
    ).to_csv(run_dir / "sc1_control_summary.csv", index=False)

    pd.DataFrame(
        [
            {
                "SCENARIO_DAIRY_COW": 1600000,
                "SCENARIO_SUCKLER_COW": 160000,
                "SCENARIO_TOTAL_CATTLE": 2500000,
                "TOTAL_CATTLE_CHANGE": -500000,
                "TOTAL_CATTLE_CHANGE_PCT": -16.6667,
            }
        ]
    ).to_csv(run_dir / "sc1_national_livestock_summary.csv", index=False)

    canonical_sc3 = _sc3_frame()
    canonical_path = run_dir / "sc3_ed_results.csv"
    canonical_sc3.to_csv(canonical_path, index=False)
    before = canonical_path.read_text(encoding="utf-8")

    outputs = export_scientific_results(run_dir)

    assert outputs["workbook"].exists()
    assert outputs["land_accounting_csv"].exists()
    assert outputs["validation_csv"].exists()
    assert canonical_path.read_text(encoding="utf-8") == before

    workbook = load_workbook(outputs["workbook"], read_only=True, data_only=True)
    assert "00_Read_Me" in workbook.sheetnames
    assert "03_SC3_Land_Accounting" in workbook.sheetnames
    assert "04_Validation" in workbook.sheetnames
    assert "10_Data_Dictionary" in workbook.sheetnames
    assert "11_ED_Results" in workbook.sheetnames
