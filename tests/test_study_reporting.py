from __future__ import annotations

from pathlib import Path
import sqlite3

import pandas as pd
from openpyxl import load_workbook

from goblin_spatial.study_reporting import (
    PRINCIPAL_RULES,
    PRINCIPAL_SCENARIOS,
    export_study_results,
)


def _write_run(root: Path, scenario: str, rule: str, scenario_i: int, rule_i: int) -> None:
    run_dir = root / f"{scenario}_2020_{rule}"
    run_dir.mkdir(parents=True)

    pd.DataFrame(
        [
            {
                "SCENARIO_NO": scenario_i + 1,
                "SCENARIO_ID": scenario,
                "SCENARIO_NAME": scenario.replace("_", " "),
                "RUN_START_YEAR": 2020,
                "TARGET_YEAR": 2050,
                "ALLOCATION_POLICY": rule,
                "PROTECTION_STRENGTH_LAMBDA": 0.5,
                "BASELINE_GRASSLAND_HA": 300.0,
                "TARGET_LIVESTOCK_LAND_HA": 100.0,
                "RUN_GROSS_RELEASE_HA": 200.0,
                "STAGE_A_TARGET_HA": 120.0,
                "GOBLIN_PARENT_AVAILABLE_BEFORE_REWETTING_HA": 80.0,
                "REWETTING_TARGET_HA": 30.0,
            }
        ]
    ).to_csv(run_dir / "sc1_control_summary.csv", index=False)

    total = 2500 - scenario_i * 200
    pd.DataFrame(
        [
            {
                "PATHWAY_NAME": scenario,
                "PATHWAY_BASELINE_YEAR": 2020,
                "MILESTONE_YEAR": 2050,
                "PATHWAY_ALLOCATION_RULE": rule,
                "PROTECTION_STRENGTH_LAMBDA": 0.5,
                "NATIONAL_COHORT_TARGET_SOURCE": "TEST",
                "SCENARIO_DAIRY_COW": 1600 - scenario_i * 200,
                "SCENARIO_SUCKLER_COW": 160 - scenario_i * 20,
                "SCENARIO_TOTAL_CATTLE": total,
                "BASE_TOTAL_CATTLE": 3000,
                "TOTAL_CATTLE_CHANGE": total - 3000,
                "TOTAL_CATTLE_CHANGE_PCT": 100.0 * (total - 3000) / 3000.0,
                "GOBLIN_RELEASED_GRASSLAND_HA": 200.0,
            }
        ]
    ).to_csv(run_dir / "sc1_national_livestock_summary.csv", index=False)

    gross_so_loss = 100000.0 + scenario_i * 20000.0 + rule_i * 5000.0
    pd.DataFrame(
        [
            {
                "SCENARIO_ID": scenario,
                "RUN_START_YEAR": 2020,
                "TARGET_YEAR": 2050,
                "ED_COUNT": 2,
                "BASE_TOTAL_CATTLE": 200.0,
                "SCENARIO_TOTAL_CATTLE": 160.0,
                "GROSS_SO_LIVESTOCK_LOSS_2020_EUR": gross_so_loss,
                "GOBLIN_RELEASED_GRASSLAND_HA": 200.0,
            }
        ]
    ).to_csv(run_dir / "sc1_national_metrics.csv", index=False)

    reductions = [20.0 + scenario_i * 3.0, 20.0 + scenario_i * 3.0]
    if rule_i == 1:
        reductions = [15.0 + scenario_i * 3.0, 25.0 + scenario_i * 3.0]
    elif rule_i == 2:
        reductions = [12.0 + scenario_i * 3.0, 28.0 + scenario_i * 3.0]
    elif rule_i == 3:
        reductions = [10.0 + scenario_i * 3.0, 30.0 + scenario_i * 3.0]

    so_exposure = [50000.0 + scenario_i * 5000.0, 50000.0 + scenario_i * 5000.0]
    if rule_i:
        so_exposure = [so_exposure[0] - rule_i * 4000.0, so_exposure[1] + rule_i * 6000.0]

    sc1 = pd.DataFrame(
        {
            "CSOED": ["1", "2"],
            "County": ["A", "B"],
            "BASE_TOTAL_CATTLE": [100.0, 100.0],
            "CUMULATIVE_REDUCTION_TOTAL_CATTLE": reductions,
            "BASE_SO_LIVESTOCK_2020_EUR": [200000.0, 200000.0],
            "SO_LIVESTOCK_EXPOSURE_2020_EUR": so_exposure,
            "GOBLIN_RELEASED_GRASSLAND_HA": [100.0, 100.0],
        }
    )
    sc1.to_csv(run_dir / "sc1_ed_results.csv", index=False)

    pd.DataFrame(
        {
            "County": ["A", "B"],
            "TOTAL_CATTLE_REDUCTION_HEAD": reductions,
            "SO_LIVESTOCK_GROSS_LOSS_2020_EUR": so_exposure,
            "GOBLIN_RELEASED_GRASSLAND_HA": [100.0, 100.0],
        }
    ).to_csv(run_dir / "sc1_county_summary.csv", index=False)

    pd.DataFrame(
        [{"CHECK": "national endpoint", "STATUS": "PASS", "DIFFERENCE": 0.0}]
    ).to_csv(run_dir / "sc1_goblin_reconciliation.csv", index=False)

    sc2 = sc1[["CSOED", "County", "GOBLIN_RELEASED_GRASSLAND_HA"]].copy()
    sc2["SC2_POTENTIAL_RELEASE_HA"] = [100.0, 100.0]
    sc2.to_csv(run_dir / "sc2_ed_context.csv", index=False)

    sc3 = sc2.copy()
    sc3["SC3_STAGE_A_REALIZED_HA"] = [60.0, 50.0]
    sc3["SC3_STAGE_A_AVAILABLE_HA"] = [40.0, 50.0]
    sc3["SC3_REALIZED_REWETTING_HA"] = [10.0, 15.0]
    sc3["SC3_RESIDUAL_AVAILABLE_LAND_HA"] = [30.0, 35.0]
    sc3.to_csv(run_dir / "sc3_ed_results.csv", index=False)

    sc3_summary = {
        "POTENTIAL_RELEASE_HA": 200.0,
        "STAGE_A_AVAILABLE_HA": 90.0,
        "REALISED_CONVERSION_HA": 135.0,
        "RESIDUAL_AVAILABLE_LAND_HA": 65.0,
        "MAX_ABS_ED_ACCOUNTING_CLOSURE_HA": 0.0,
        "TOTAL_UNMET_TARGET_HA": 15.0,
        "ACCOUNTING_CLOSURE_HA": 0.0,
    }
    values = {
        "AD_GRASS": (40.0, 40.0, 0.0),
        "BIOREFINERY_GRASS": (20.0, 20.0, 0.0),
        "WILLOW": (20.0, 20.0, 0.0),
        "ADDITIONAL_TILLAGE": (20.0, 10.0, 10.0),
        "FOREST": (20.0, 20.0, 0.0),
        "REWETTING": (30.0, 25.0, 5.0),
    }
    for use, (target, realised, unmet) in values.items():
        sc3_summary[f"TARGET_{use}_HA"] = target
        sc3_summary[f"REALISED_{use}_HA"] = realised
        sc3_summary[f"UNMET_{use}_HA"] = unmet
    pd.DataFrame([sc3_summary]).to_csv(run_dir / "sc3_national_summary.csv", index=False)


def test_final_study_reporting_builds_workbook_sqlite_and_figures(tmp_path: Path) -> None:
    root = tmp_path / "runs"
    for scenario_i, scenario in enumerate(PRINCIPAL_SCENARIOS):
        for rule_i, rule in enumerate(PRINCIPAL_RULES):
            _write_run(root, scenario, rule, scenario_i, rule_i)

    out = tmp_path / "final"
    outputs = export_study_results(root, output_dir=out)

    assert outputs["workbook"].exists()
    assert outputs["sqlite"].exists()
    assert outputs["figure_data_csv"].exists()
    assert outputs["figure_directory"].exists()
    assert (outputs["figure_directory"] / "fig01_livestock_endpoint.png").exists()
    assert (outputs["figure_directory"] / "fig01_livestock_endpoint.svg").exists()
    assert (outputs["figure_directory"] / "fig03_released_land_allocation.png").exists()

    workbook = load_workbook(outputs["workbook"], read_only=True, data_only=True)
    for sheet in (
        "00_Read_Me",
        "02_National_Results",
        "05_Land_Accounting",
        "08_Protection_PRORATA",
        "10_Robust_Exposure",
        "16_Figure_Data",
        "18_Figures",
    ):
        assert sheet in workbook.sheetnames

    with sqlite3.connect(outputs["sqlite"]) as connection:
        run_count = connection.execute("SELECT COUNT(*) FROM run_registry").fetchone()[0]
        robust_count = connection.execute("SELECT COUNT(*) FROM robust_exposure").fetchone()[0]
        figure_count = connection.execute("SELECT COUNT(*) FROM figure_data").fetchone()[0]
    assert run_count == 12
    assert robust_count == 2
    assert figure_count > 0
